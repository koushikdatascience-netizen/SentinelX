"""Persisted visual batches and explicit, immutable campaign approval snapshots."""
import csv
import hashlib
import io
import json
import zipfile
from typing import Literal
from uuid import UUID,uuid5
from fastapi import APIRouter,Depends,HTTPException,Response
from pydantic import BaseModel,Field
from sqlalchemy import select,update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.auth import current_user
from app.workspace_store import User,WorkflowRun,CampaignProject,CampaignItem,CampaignCreative,ApprovalRecord,ApprovalEvent,session,utcnow
from app.approval_routes import authorized_campaign,ensure_board
from app.campaign_routes import generation_is_live
from app.creatives import RequestVariants,generate_variants,_variant_dict

router=APIRouter(tags=["Campaign delivery"])


def digest(value):return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False).encode()).hexdigest()


def batch_info(run):return {"id":run.id,"status":run.status,**json.loads(run.result_json)}


def batch_owned(db,user,cid,wid,bid,lock=False):
    authorized_campaign(db,cid,wid,user)
    query=select(WorkflowRun).where(WorkflowRun.id==bid,WorkflowRun.workspace_id==wid,WorkflowRun.website=="campaign-visual-batch:"+cid)
    run=db.scalar(query.with_for_update() if lock else query)
    if not run:raise HTTPException(404,"Visual batch not found")
    return run


class BatchRequest(BaseModel):
    request_id:UUID
    generation_mode:Literal["live","sample"]="live"
    select_drafts:bool=False


@router.get("/api/v1/campaigns/{cid}/creative-status")
def creative_status(cid:str,workspace_id:str,user:User=Depends(current_user),db:Session=Depends(session)):
    authorized_campaign(db,cid,workspace_id,user)
    variants=db.scalars(select(CampaignCreative).where(CampaignCreative.campaign_id==cid).order_by(CampaignCreative.created_at)).all()
    return [_variant_dict(v) for v in variants]


@router.post("/api/v1/campaigns/{cid}/visual-batches",status_code=201)
def start_batch(cid:str,workspace_id:str,body:BatchRequest,user:User=Depends(current_user),db:Session=Depends(session)):
    campaign=authorized_campaign(db,cid,workspace_id,user)
    db.scalar(select(CampaignProject).where(CampaignProject.id==cid,CampaignProject.workspace_id==workspace_id).with_for_update())
    generation_is_live(user,body.generation_mode)
    items=db.scalars(select(CampaignItem).where(CampaignItem.campaign_id==cid).order_by(CampaignItem.day)).all()
    signature=digest({"campaign":cid,"items":[{"id":item.id,"content":item.item_json} for item in items],"mode":body.generation_mode,"select":body.select_drafts})
    old=db.get(WorkflowRun,str(body.request_id))
    if old:
        data=json.loads(old.result_json)
        if old.workspace_id!=workspace_id:raise HTTPException(404,"Batch request not found")
        if data.get("operation")!="campaign_visual_batch" or data.get("signature")!=signature:raise HTTPException(409,"Request ID belongs to different batch inputs")
        return batch_info(old)
    active=db.scalar(select(WorkflowRun).where(WorkflowRun.workspace_id==workspace_id,WorkflowRun.website=="campaign-visual-batch:"+cid,WorkflowRun.status.in_(["pending","running"])))
    if active:raise HTTPException(409,"A visual batch is already pending or running. Open its saved status before starting another.")
    result={"operation":"campaign_visual_batch","signature":signature,"generation_mode":body.generation_mode,"select_drafts":body.select_drafts,
        "items":[{"id":item.id,"day":item.day,"content_hash":hashlib.sha256(item.item_json.encode()).hexdigest()} for item in items],"completed":[],"total":len(items),"error":None}
    run=WorkflowRun(id=str(body.request_id),workspace_id=workspace_id,website="campaign-visual-batch:"+cid,objective=campaign.goal,status="pending",result_json=json.dumps(result))
    db.add(run)
    try:db.commit()
    except IntegrityError:db.rollback();raise HTTPException(409,"Batch request already exists; no duplicate generation started") from None
    return batch_info(run)


@router.get("/api/v1/campaigns/{cid}/visual-batches")
def latest_batch(cid:str,workspace_id:str,user:User=Depends(current_user),db:Session=Depends(session)):
    authorized_campaign(db,cid,workspace_id,user)
    run=db.scalar(select(WorkflowRun).where(WorkflowRun.workspace_id==workspace_id,WorkflowRun.website=="campaign-visual-batch:"+cid).order_by(WorkflowRun.created_at.desc()))
    return batch_info(run) if run else None


@router.post("/api/v1/campaigns/{cid}/visual-batches/{bid}/next")
def next_visual(cid:str,bid:str,workspace_id:str,user:User=Depends(current_user),db:Session=Depends(session)):
    run=batch_owned(db,user,cid,workspace_id,bid,True);data=json.loads(run.result_json)
    if run.status in ("completed","failed"):return batch_info(run)
    if run.status=="running":raise HTTPException(409,"The previous image request is still running or was interrupted. No automatic paid replay.")
    changed=db.execute(update(WorkflowRun).where(WorkflowRun.id==bid,WorkflowRun.status=="pending").values(status="running"))
    if changed.rowcount!=1:db.rollback();raise HTTPException(409,"This batch step is already claimed")
    db.commit();run.status="running"
    try:
        generation_is_live(user,data["generation_mode"])
        index=len(data["completed"])
        if index<len(data["items"]):
            entry=data["items"][index];item=db.get(CampaignItem,entry["id"])
            if not item or item.campaign_id!=cid or hashlib.sha256(item.item_json.encode()).hexdigest()!=entry["content_hash"]:
                raise ValueError("Campaign content changed during this batch. Start a new reviewed batch.")
            variants=db.scalars(select(CampaignCreative).where(CampaignCreative.campaign_id==cid,CampaignCreative.item_id==item.id).order_by(CampaignCreative.created_at)).all()
            if not variants:
                request=RequestVariants(variants=1,evaluate=False,generation_mode=data["generation_mode"],request_id=uuid5(UUID(bid),item.id))
                response=generate_variants(cid,item.id,workspace_id,request,user,db,budget_run_id=bid)
                if not response["created_count"]:raise ValueError("Image generation stopped. Saved images are preserved; no automatic retry. Review the batch and workspace budget before a new request.")
                variants=db.scalars(select(CampaignCreative).where(CampaignCreative.campaign_id==cid,CampaignCreative.item_id==item.id).order_by(CampaignCreative.created_at)).all()
            if data["select_drafts"] and not any(v.selected for v in variants):variants[0].selected=True
            data["completed"].append({"item_id":item.id,"day":entry["day"],"variant_ids":[v.id for v in variants]})
        run.status="completed" if len(data["completed"])==len(data["items"]) else "pending"
    except Exception as exc:
        db.rollback();run.status="failed"
        data["error"]=str(exc) if isinstance(exc,ValueError) and str(exc).startswith(("Campaign content changed","Image generation stopped")) else "Visual batch stopped. Existing images are saved; no automatic paid retry. Check live settings and app/workspace/run budget."
    run.result_json=json.dumps(data);db.commit()
    return batch_info(run)


def manifest(db,campaign):
    posts=[]
    items=db.scalars(select(CampaignItem).where(CampaignItem.campaign_id==campaign.id).order_by(CampaignItem.day)).all()
    for item in items:
        visuals=db.scalars(select(CampaignCreative).where(CampaignCreative.campaign_id==campaign.id,CampaignCreative.item_id==item.id,CampaignCreative.selected.is_(True))).all()
        if len(visuals)!=1:raise HTTPException(409,f"Choose exactly one saved visual for day {item.day} before final approval")
        visual=visuals[0];content=json.loads(item.item_json)
        posts.append({"item_id":item.id,"day":item.day,"platform":item.platform,"state":item.state,"provenance":item.provenance,
            **{key:content.get(key) for key in ("post_type","caption","hashtags","cta","creative_brief","image_prompt","video_script","evidence_ids","unsupported_claims")},
            "visual":{"id":visual.id,"mime":visual.mime,"provenance":visual.provenance,"generation_prompt":visual.prompt,
                "sha256":hashlib.sha256(visual.image_blob).hexdigest(),"file":"creatives/"+visual.id+"."+{"image/png":"png","image/jpeg":"jpg","image/webp":"webp"}.get(visual.mime,"bin"),
                "image_url":f"/api/v1/campaigns/{campaign.id}/items/{item.id}/creatives/{visual.id}/image?workspace_id={campaign.workspace_id}"}})
    if not posts:raise HTTPException(409,"Campaign has no posts")
    return {"campaign_id":campaign.id,"workspace_id":campaign.workspace_id,"goal":campaign.goal,"language":campaign.language,"provenance":campaign.provenance,"posts":posts,"externally_published":False}


def snapshot_current(db,campaign,run):
    data=json.loads(run.result_json)
    current=manifest(db,campaign)
    if digest(current)!=data.get("content_hash"):return False
    expected={("campaign_item",p["item_id"]) for p in current["posts"]}|{("creative",p["visual"]["id"]) for p in current["posts"]}
    approved={(r.target_kind,r.target_id) for r in db.scalars(select(ApprovalRecord).where(ApprovalRecord.campaign_id==campaign.id,ApprovalRecord.state=="approved")).all()}
    return expected<=approved


class FinalReview(BaseModel):
    request_id:UUID
    reviewed:bool=False
    note:str=Field(default="",max_length=500)


@router.post("/api/v1/campaigns/{cid}/finalize",status_code=201)
def finalize(cid:str,workspace_id:str,body:FinalReview,user:User=Depends(current_user),db:Session=Depends(session)):
    campaign=authorized_campaign(db,cid,workspace_id,user)
    if not body.reviewed:raise HTTPException(422,"Explicit confirmation of the complete posts and selected visuals is required")
    old=db.get(WorkflowRun,str(body.request_id))
    if old:
        if old.workspace_id!=workspace_id or old.website!="campaign-final:"+cid:raise HTTPException(404,"Final approval not found")
        if not snapshot_current(db,campaign,old):raise HTTPException(409,"This approval snapshot is no longer current")
        return {"id":old.id,"current":True,**json.loads(old.result_json)}
    manifest(db,campaign)  # Validate completeness before recording any decisions.
    ensure_board(db,cid,workspace_id)
    current=manifest(db,campaign)
    targets={("campaign_item",p["item_id"]) for p in current["posts"]}|{("creative",p["visual"]["id"]) for p in current["posts"]}
    records=db.scalars(select(ApprovalRecord).where(ApprovalRecord.campaign_id==cid).with_for_update()).all()
    if any(r.state=="rejected" for r in records if (r.target_kind,r.target_id) in targets):raise HTTPException(409,"Reopen rejected items explicitly before final approval")
    for record in records:
        if (record.target_kind,record.target_id) not in targets or record.state=="approved":continue
        record.state="approved";db.add(ApprovalEvent(approval_id=record.id,reviewer_id=user.id,decision="approved",note=body.note or "Final campaign review"))
        if record.target_kind=="campaign_item":db.get(CampaignItem,record.target_id).state="approved"
    db.flush();current=manifest(db,campaign)
    result={"operation":"campaign_final_approval","approved_by":user.id,"approved_at":utcnow().isoformat(),"content_hash":digest(current),"manifest":current}
    run=WorkflowRun(id=str(body.request_id),workspace_id=workspace_id,website="campaign-final:"+cid,objective=campaign.goal,status="completed",result_json=json.dumps(result,ensure_ascii=False))
    db.add(run)
    try:db.commit()
    except IntegrityError:db.rollback();raise HTTPException(409,"Final approval already exists; refresh its saved status") from None
    return {"id":run.id,"current":True,**result}


@router.get("/api/v1/campaigns/{cid}/finalized")
def finalized(cid:str,workspace_id:str,user:User=Depends(current_user),db:Session=Depends(session)):
    campaign=authorized_campaign(db,cid,workspace_id,user)
    run=db.scalar(select(WorkflowRun).where(WorkflowRun.workspace_id==workspace_id,WorkflowRun.website=="campaign-final:"+cid).order_by(WorkflowRun.created_at.desc()))
    if not run:return None
    try:valid=snapshot_current(db,campaign,run)
    except HTTPException:valid=False
    return {"id":run.id,"current":valid,**json.loads(run.result_json)}


def spreadsheet_rows(data):
    columns=["day","platform","post_type","caption","hashtags","cta","creative_brief","image_prompt","video_script","provenance","visual_file","visual_provenance","visual_sha256","image_url","generation_prompt"]
    rows=[]
    for post in data["posts"]:
        item={**post,"hashtags":" ".join(post.get("hashtags") or []),**{"visual_file":post["visual"]["file"],"visual_provenance":post["visual"]["provenance"],"visual_sha256":post["visual"]["sha256"],"image_url":post["visual"]["image_url"],"generation_prompt":post["visual"]["generation_prompt"]}}
        rows.append([item.get(key) or "" for key in columns])
    return columns,rows


def csv_bytes(data):
    columns,rows=spreadsheet_rows(data);out=io.StringIO(newline="");writer=csv.writer(out);writer.writerow(columns)
    def safe(value):
        return "'"+value if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")) else value
    writer.writerows([[safe(value) for value in row] for row in rows]);return out.getvalue().encode("utf-8-sig")


def excel_bytes(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    columns,rows=spreadsheet_rows(data);book=Workbook();sheet=book.active;sheet.title="Approved posts";sheet.append(columns)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type="s"
            cell.alignment=Alignment(vertical="top",wrap_text=True)
    for cell in sheet[1]:cell.font=Font(color="FFFFFF",bold=True);cell.fill=PatternFill("solid",fgColor="354D83")
    sheet.freeze_panes="D2";sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=18 if column[0].value in ("day","platform","provenance") else 42
    for i in range(2,sheet.max_row+1):sheet.row_dimensions[i].height=80
    out=io.BytesIO();book.save(out);return out.getvalue()


@router.get("/api/v1/campaigns/{cid}/finalized/{rid}/export")
def export_final(cid:str,rid:str,workspace_id:str,format:Literal["json","csv","xlsx","zip"]="json",user:User=Depends(current_user),db:Session=Depends(session)):
    campaign=authorized_campaign(db,cid,workspace_id,user)
    run=db.scalar(select(WorkflowRun).where(WorkflowRun.id==rid,WorkflowRun.workspace_id==workspace_id,WorkflowRun.website=="campaign-final:"+cid))
    if not run:raise HTTPException(404,"Final campaign not found")
    if not snapshot_current(db,campaign,run):raise HTTPException(409,"Campaign or approvals changed; complete a new final review before exporting")
    saved=json.loads(run.result_json);data=saved["manifest"]
    payload=json.dumps(saved,ensure_ascii=False,indent=2).encode();mime="application/json"
    if format=="csv":payload=csv_bytes(data);mime="text/csv"
    elif format=="xlsx":payload=excel_bytes(data);mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif format=="zip":
        out=io.BytesIO()
        with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("campaign.json",payload);archive.writestr("posts.csv",csv_bytes(data));archive.writestr("posts.xlsx",excel_bytes(data))
            for post in data["posts"]:archive.writestr(post["visual"]["file"],db.get(CampaignCreative,post["visual"]["id"]).image_blob)
            archive.writestr("README.txt","Owner-approved campaign snapshot. Nothing was published externally. Images are inside creatives/. API image links require workspace authentication.")
        payload=out.getvalue();mime="application/zip"
    return Response(payload,media_type=mime,headers={"Content-Disposition":f'attachment; filename="sentinelx-final-{cid}.{format}"',"Cache-Control":"private, no-store","X-Content-Type-Options":"nosniff"})
