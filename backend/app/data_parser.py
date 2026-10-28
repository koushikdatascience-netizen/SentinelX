"""Bounded deterministic file interpretation. Run untrusted parsers in a subprocess."""
import csv
import io
import json
import re
import sys
import zipfile
from datetime import datetime,date
from decimal import Decimal,InvalidOperation

MAX_ROWS=2000
MAX_COLUMNS=40
MAX_BYTES=2_000_000


def decode(payload):
    if payload.startswith((b"\xff\xfe",b"\xfe\xff")):return payload.decode("utf-16"),"UTF-16 BOM"
    try:return payload.decode("utf-8-sig"),"UTF-8"
    except UnicodeDecodeError:raise ValueError("Use UTF-8 CSV/TXT/JSON or UTF-16 with a byte-order mark") from None


def cell(value):
    if value is None:return ""
    if isinstance(value,(date,datetime)):return value.isoformat()
    text=str(value).strip()
    if len(text)>4000:raise ValueError("A cell exceeds 4,000 characters")
    return text


def number(value):
    text=value.strip();unit="number"
    for prefix,name in (("₹","INR"),("INR ","INR"),("USD ","USD"),("$","dollar_unspecified"),("€","EUR"),("£","GBP")):
        if text.startswith(prefix):text=text[len(prefix):].strip();unit=name;break
    if "," in text:
        if not re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?",text):return None
        text=text.replace(",","")
    if not re.fullmatch(r"-?\d{1,15}(\.\d{1,6})?",text):return None
    try:return Decimal(text),unit
    except InvalidOperation:return None


def tabular(rows):
    if not rows:raise ValueError("No rows were found")
    headers=[cell(value) for value in rows[0]]
    if not 1<=len(headers)<=MAX_COLUMNS or not all(headers) or len(set(headers))!=len(headers):
        raise ValueError("Use 1–40 unique nonempty column headings")
    values=[];characters=0
    for row in rows[1:]:
        if len(values)>=MAX_ROWS:raise ValueError("Dataset exceeds 2,000 rows; upload a bounded export")
        if len(row)>len(headers):raise ValueError("Rows have more values than the header")
        converted=[cell(value) for value in row]+[""]*(len(headers)-len(row))
        characters+=sum(len(value) for value in converted)
        if characters>1_000_000:raise ValueError("Extracted table exceeds one million characters")
        values.append(converted)
    columns=[]
    for index,header in enumerate(headers):
        present=[row[index] for row in values if row[index]]
        parsed=[number(value) for value in present]
        units={value[1] for value in parsed if value}
        numeric=bool(present) and all(value is not None for value in parsed) and len(units)==1
        dates=bool(present) and all(re.fullmatch(r"\d{4}-\d{2}-\d{2}(T.*)?",value) for value in present)
        if dates:
            try:
                for value in present:datetime.fromisoformat(value)
            except ValueError:dates=False
        details={"name":header,"type":"number" if numeric else "iso_date" if dates else "text",
            "missing":len(values)-len(present),"distinct":len(set(present)),
            "description":"Detected from supplied values; business meaning and units require owner review."}
        if numeric:
            amounts=[value[0] for value in parsed]
            details.update(unit=next(iter(units)),sum=str(sum(amounts,Decimal(0))),minimum=str(min(amounts)),maximum=str(max(amounts)))
        if len(units)>1:details["warning"]="Mixed units; no total computed"
        columns.append(details)
    return {"kind":"table","columns":columns,"rows":values,"row_count":len(values),
        "duplicate_rows":len(values)-len(set(tuple(row) for row in values)),
        "formula_like_cells":sum(value.lstrip().startswith(("=","+","-","@")) and number(value) is None for row in values for value in row),
        "warnings":["No formulas executed; dates are recognized only in unambiguous ISO form. No currency conversion or cross-file join performed."]}


def parse_file(payload,extension):
    if not 0<len(payload)<=MAX_BYTES:raise ValueError("Each file must be between 1 byte and 2 MB")
    if extension in ("csv","txt","json"):
        text,encoding=decode(payload)
        if "\0" in text:raise ValueError("Binary content is not accepted as text")
        if extension=="txt":return {"kind":"document","text":text[:50000],"characters":len(text),"truncated":len(text)>50000,"encoding":encoding}
        if extension=="json":
            data=json.loads(text)
            if not isinstance(data,list) or not data or len(data)>MAX_ROWS or not all(isinstance(row,dict) for row in data):
                raise ValueError("JSON must be an array of 1–2,000 flat row objects")
            headers=list(data[0])
            if any(set(row)!=set(headers) for row in data):raise ValueError("JSON rows require the same fields")
            if any(isinstance(value,(dict,list)) for row in data for value in row.values()):raise ValueError("Nested JSON rows are unsupported")
            result=tabular([headers]+[[row[name] for name in headers] for row in data])
        else:
            try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=",;\t|")
            except csv.Error:dialect=csv.excel
            reader=csv.reader(io.StringIO(text,newline=""),dialect=dialect,strict=True)
            rows=[]
            for row in reader:
                if len(rows)>MAX_ROWS:raise ValueError("CSV exceeds 2,000 data rows")
                if row:rows.append(row)
            result=tabular(rows)
        return {**result,"encoding":encoding}
    if extension=="xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            infos=archive.infolist()
            if len(infos)>1000 or sum(item.file_size for item in infos)>20_000_000:raise ValueError("Expanded workbook exceeds safe limits")
            if any(item.file_size>10_000_000 or item.file_size>max(1,item.compress_size)*200 for item in infos):raise ValueError("Workbook compression exceeds safe limits")
            if any("vbaproject" in item.filename.lower() or "externallinks" in item.filename.lower() for item in infos):raise ValueError("Macros and external workbook links are unsupported")
        book=load_workbook(io.BytesIO(payload),read_only=True,data_only=False,keep_links=False)
        try:
            if len(book.worksheets)!=1:raise ValueError("Export one worksheet per file so schemas are reviewed separately")
            sheet=book.worksheets[0]
            if sheet.max_row and sheet.max_row>MAX_ROWS+1 or sheet.max_column and sheet.max_column>MAX_COLUMNS:raise ValueError("Workbook exceeds 2,000 rows or 40 columns")
            rows=[]
            for row in sheet.iter_rows(values_only=True):
                if len(rows)>MAX_ROWS:raise ValueError("Workbook exceeds 2,000 data rows")
                rows.append(list(row))
            result=tabular(rows)
            result["worksheet"]=sheet.title;return result
        finally:book.close()
    if extension=="pdf":
        from pypdf import PdfReader
        if not payload.startswith(b"%PDF-"):raise ValueError("Invalid PDF signature")
        reader=PdfReader(io.BytesIO(payload),strict=True)
        if reader.is_encrypted:raise ValueError("Encrypted PDFs are unsupported")
        if len(reader.pages)>20:raise ValueError("PDF exceeds 20 pages")
        pages=[]
        for index,page in enumerate(reader.pages):
            contents=page.get_contents()
            if contents is not None and len(contents.get_data())>5_000_000:raise ValueError("PDF page stream exceeds safe limits")
            pages.append({"page":index+1,"text":(page.extract_text() or "")[:4000]})
        text="\n".join(page["text"] for page in pages)
        if not text.strip():raise ValueError("No extractable text; scanned PDFs require OCR, which is not enabled")
        return {"kind":"document","pages":pages,"text":text[:50000],"truncated":len(text)>50000,
            "warnings":["Text extraction only; images, scripts and links were not executed. Page text is bounded to 4,000 characters."]}
    raise ValueError("Supported file types: CSV, XLSX, PDF, TXT, JSON")


def main():
    # Unix worker limits apply before importing document libraries. No network,
    # shell execution, user-selected paths or environment secrets are needed.
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS,(512*1024*1024,512*1024*1024))
        resource.setrlimit(resource.RLIMIT_CPU,(12,12))
        resource.setrlimit(resource.RLIMIT_FSIZE,(0,0))
    except ImportError:pass  # Windows local development still has wall/size limits.
    try:
        payload=sys.stdin.buffer.read(MAX_BYTES+1)
        result=parse_file(payload,sys.argv[1])
        print(json.dumps({"ok":True,"data":result},ensure_ascii=True,allow_nan=False))
    except Exception as exc:
        print(json.dumps({"ok":False,"error":str(exc)[:200] if isinstance(exc,ValueError) else "File parsing failed ("+type(exc).__name__+")"}))


if __name__=="__main__":main()
