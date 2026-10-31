"""Bounded, local CSV/XLSX calculations. Never executes model-generated code."""
import csv
import io
import json
import zipfile
from collections import defaultdict
from datetime import date,datetime
from decimal import Decimal,InvalidOperation
from typing import Literal
from pydantic import BaseModel,Field,ConfigDict

MAX_ROWS=20000


def numeric(value):
    if isinstance(value,bool) or value is None:return None
    if isinstance(value,(int,float,Decimal)):text=str(value)
    else:text=str(value).strip().replace(",","").removeprefix("₹").removeprefix("$").strip()
    if len(text)>200:return None
    if text.startswith("(") and text.endswith(")"):text="-"+text[1:-1]
    try:
        number=Decimal(text)
        return number if number.is_finite() and (not number or -100<=number.adjusted()<=100) else None
    except (InvalidOperation,ValueError):return None


def table(payload,filename,sheet=None):
    if not 0<len(payload)<=8_000_000:raise ValueError("Datasets must be nonempty and at most 8 MB")
    names=[];book=None
    if filename.lower().endswith(".csv"):
        try:text=payload.decode("utf-16" if payload.startswith((b"\xff\xfe",b"\xfe\xff")) else "utf-8-sig")
        except UnicodeError:raise ValueError("Save CSV as UTF-8 or UTF-16 with a byte-order mark") from None
        try:dialect=csv.Sniffer().sniff(text[:4000],delimiters=",;\t")
        except csv.Error:dialect=csv.excel
        iterator=iter(csv.reader(io.StringIO(text),dialect))
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                entries=archive.infolist()
                if len(entries)>1000 or sum(entry.file_size for entry in entries)>20_000_000 or any("vbaproject" in entry.filename.lower() for entry in entries):raise ValueError("Workbook exceeds safe expansion limits or contains macros")
            book=load_workbook(io.BytesIO(payload),read_only=True,data_only=True,keep_links=False)
            names=book.sheetnames[:30]
            if sheet and sheet not in names:raise ValueError("Choose a worksheet present in this workbook")
            selected=book[sheet or names[0]];iterator=iter(selected.iter_rows(max_row=MAX_ROWS+2,values_only=True))
        except (OSError,zipfile.BadZipFile,KeyError):raise ValueError("The XLSX workbook could not be read") from None
    else:raise ValueError("Use CSV or XLSX for local data analysis")
    try:
        header=next(iterator,None)
        if not header:raise ValueError("Dataset has no header row")
        columns=[str(value).strip() if value is not None else "" for value in header]
        if not 1<=len(columns)<=60 or any(not column or len(column)>100 for column in columns) or len(set(columns))!=len(columns):raise ValueError("Use 1 to 60 unique, nonempty column headings, each up to 100 characters")
        rows=[];truncated=False
        for index,values in enumerate(iterator):
            if index>=MAX_ROWS:truncated=True;break
            if not any(value is not None and str(value).strip() for value in values):continue
            if len(values)>len(columns) and any(value is not None and str(value).strip() for value in values[len(columns):]):raise ValueError("A data row contains more fields than the header")
            rows.append({column:(value.isoformat() if isinstance(value,(datetime,date)) else value) for column,value in zip(columns,list(values)+[None]*len(columns))})
        return {"columns":columns,"rows":rows,"truncated":truncated,"sheet_names":names,"sheet":sheet or (names[0] if names else None)}
    except csv.Error:raise ValueError("The CSV contains malformed or oversized fields") from None
    finally:
        if book:book.close()


def profile(payload,filename):
    data=table(payload,filename);stats=[]
    for column in data["columns"]:
        values=[row[column] for row in data["rows"] if row[column] is not None and str(row[column]).strip()]
        numbers=[numeric(value) for value in values];usable=[number for number in numbers if number is not None]
        is_number=bool(values) and len(usable)>=len(values)*0.8
        item={"name":column,"nonempty":len(values),"missing":len(data["rows"])-len(values),"numeric":is_number}
        if is_number:item.update(sum=str(sum(usable,Decimal(0))),minimum=str(min(usable)),maximum=str(max(usable)),average=str(sum(usable,Decimal(0))/len(usable)),nonnumeric_ignored=len(values)-len(usable))
        stats.append(item)
    return {"rows_analyzed":len(data["rows"]),"columns":stats,"truncated":data["truncated"],"sheet":data["sheet"],"sheet_names":data["sheet_names"],"provenance":"user_supplied","method":"local_calculation","limitation":"Up to 20000 physical data rows and 60 columns. Numeric-looking identifiers may not be meaningful measures. Workbook values are used as saved; formulas are not recalculated. No growth cause or future result is established."}


class DataFilter(BaseModel):
    model_config=ConfigDict(extra="forbid")
    column:str=Field(min_length=1,max_length=100)
    operator:Literal["eq","contains","gt","gte","lt","lte"]="eq"
    value:str=Field(max_length=200)


class DataQuery(BaseModel):
    model_config=ConfigDict(extra="forbid")
    limit:int=Field(default=20,ge=1,le=20)
    aggregation:Literal["sum","average","minimum","maximum","count","distinct_count"]="sum"
    metric:str|None=Field(default=None,max_length=100)
    group_by:str|None=Field(default=None,max_length=100)
    group_interval:Literal["value","month","year"]="value"
    filters:list[DataFilter]=Field(default_factory=list,max_length=5)
    sort:Literal["value_desc","group_asc"]="value_desc"
    sheet:str|None=Field(default=None,max_length=100)


def calculate(payload,filename,query):
    data=table(payload,filename,query.sheet);columns=set(data["columns"])
    requested=[query.metric,query.group_by,*[entry.column for entry in query.filters]]
    if any(column and column not in columns for column in requested):raise ValueError("The requested column is not in this dataset")
    if query.aggregation!="count" and not query.metric:raise ValueError("Choose a metric column for this calculation")
    rows=data["rows"]
    for entry in query.filters:
        def match(row):
            value=row[entry.column]
            if entry.operator=="eq":return str("" if value is None else value).casefold()==entry.value.casefold()
            if entry.operator=="contains":return entry.value.casefold() in str("" if value is None else value).casefold()
            left=numeric(value);right=numeric(entry.value)
            if left is None or right is None:return False
            return {"gt":left>right,"gte":left>=right,"lt":left<right,"lte":left<=right}[entry.operator]
        rows=[row for row in rows if match(row)]
    groups=defaultdict(list)
    for row in rows:
        group=str(row.get(query.group_by) if row.get(query.group_by) is not None else "(missing)") if query.group_by else "All matching rows"
        if query.group_by and query.group_interval!="value":
            try:
                day=datetime.fromisoformat(str(row[query.group_by])[:10]);group=day.strftime("%Y-%m" if query.group_interval=="month" else "%Y")
            except (ValueError,TypeError):group="(unparseable date)"
        groups[group].append(row.get(query.metric) if query.metric else 1)
    results=[];ignored=0
    for group,values in groups.items():
        numbers=[numeric(value) for value in values];valid=[number for number in numbers if number is not None]
        if query.aggregation=="count":value=Decimal(len(values))
        elif query.aggregation=="distinct_count":value=Decimal(len({str(value) for value in values if value is not None}))
        else:
            ignored+=len(values)-len(valid)
            if not valid:continue
            value={"sum":lambda:sum(valid,Decimal(0)),"average":lambda:sum(valid,Decimal(0))/len(valid),"minimum":lambda:min(valid),"maximum":lambda:max(valid)}[query.aggregation]()
        results.append({"group":group,"value":str(value)})
    results.sort(key=(lambda row:row["group"]) if query.sort=="group_asc" else (lambda row:-Decimal(row["value"])))
    return {"method":"local_calculation","provenance":"user_supplied","query":query.model_dump(),"rows_analyzed":len(data["rows"]),"matching_rows":len(rows),"results":results[:query.limit],"groups_total":len(results),"nonnumeric_ignored":ignored,"truncated_input":data["truncated"],"limitation":"Computed locally from uploaded rows, not AI-estimated. At most 20 result groups shown. Workbook formulas are not recalculated; source data and business causation remain unverified."}
