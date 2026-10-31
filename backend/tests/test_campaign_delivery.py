import base64
import csv
import io
import json
import zipfile
from uuid import uuid4
import httpx
import pytest
from openai import OpenAI as SDK
from openpyxl import load_workbook
from app import creatives,meta_client
from app.workspace_store import SessionLocal,WorkflowRun,CampaignCreative
from tests.test_meta_client import png
from tests.test_business_visuals import enable


def campaign(client):
    wid=client.post("/api/v1/workspaces",json={"name":"Campaign archive"}).json()["id"]
    response=client.post("/api/v1/campaigns",json={"workspace_id":wid,"business_description":"Custom wood components made to supplied drawings for trade buyers.","goal":"Qualified trade inquiries for custom wooden parts","platforms":["instagram"],"generation_mode":"sample"})
    assert response.status_code==201,response.text
    return wid,response.json()


def prepare(client,monkeypatch):
    monkeypatch.setattr(creatives,"sample_gradient",lambda seed:png())
    wid,plan=campaign(client);base=f"/api/v1/campaigns/{plan['id']}"
    body={"request_id":str(uuid4()),"generation_mode":"sample","select_drafts":True}
    first=client.post(base+"/visual-batches",params={"workspace_id":wid},json=body)
    assert first.status_code==201,first.text
    assert client.post(base+"/visual-batches",params={"workspace_id":wid},json=body).json()["id"]==first.json()["id"]
    bid=first.json()["id"]
    for _ in range(7):
        response=client.post(base+f"/visual-batches/{bid}/next",params={"workspace_id":wid})
        assert response.status_code==200,response.text
        assert response.json()["status"]!="failed",response.text
    assert response.json()["status"]=="completed"
    return wid,plan,base,bid


def test_campaign_history_preserves_earlier_campaigns(authenticated_client):
    wid,old=campaign(authenticated_client)
    new=authenticated_client.post("/api/v1/campaigns",json={"workspace_id":wid,"business_description":"A second campaign for the same workshop and trade audience.","goal":old["goal"],"generation_mode":"sample"}).json()
    saved=authenticated_client.get("/api/v1/campaigns",params={"workspace_id":wid}).json()
    assert {old["id"],new["id"]}<={row["id"] for row in saved}
    assert all(row["created_at"] and row["post_count"]==7 for row in saved)
    assert authenticated_client.get(f"/api/v1/campaigns/{old['id']}",params={"workspace_id":wid}).json()["items"]==old["items"]


def test_bulk_visuals_are_persisted_reused_and_tenant_isolated(authenticated_client,monkeypatch):
    wid,plan,base,bid=prepare(authenticated_client,monkeypatch)
    status=authenticated_client.get(base+"/creative-status",params={"workspace_id":wid}).json()
    assert len(status)==7 and all(v["selected"] for v in status)
    assert authenticated_client.post(base+f"/visual-batches/{bid}/next",params={"workspace_id":wid}).json()["status"]=="completed"
    again=authenticated_client.post(base+"/visual-batches",params={"workspace_id":wid},json={"request_id":str(uuid4()),"generation_mode":"sample","select_drafts":True}).json()
    for _ in range(7):authenticated_client.post(base+f"/visual-batches/{again['id']}/next",params={"workspace_id":wid})
    assert len(authenticated_client.get(base+"/creative-status",params={"workspace_id":wid}).json())==7
    assert authenticated_client.get(base+"/visual-batches",params={"workspace_id":"foreign"}).status_code==404
    assert authenticated_client.post(base+f"/visual-batches/{bid}/next",params={"workspace_id":"foreign"}).status_code==404


def test_final_review_is_complete_immutable_exportable_and_revocable(authenticated_client,monkeypatch):
    wid,plan,base,bid=prepare(authenticated_client,monkeypatch)
    iid=plan["items"][0]["id"]
    # CSV must neutralize formulas; JSON preserves the exact owner-approved text.
    caption="=HYPERLINK(\"https://example.org\",\"A saved owner caption\")"
    assert authenticated_client.patch(base+f"/items/{iid}",params={"workspace_id":wid},json={"caption":caption}).status_code==200
    body={"request_id":str(uuid4()),"reviewed":True}
    assert authenticated_client.post(base+"/finalize",params={"workspace_id":wid},json={**body,"reviewed":False}).status_code==422
    final=authenticated_client.post(base+"/finalize",params={"workspace_id":wid},json=body)
    assert final.status_code==201,final.text
    assert authenticated_client.post(base+"/finalize",params={"workspace_id":wid},json=body).json()["id"]==final.json()["id"]
    path=base+f"/finalized/{final.json()['id']}/export"
    raw=authenticated_client.get(path,params={"workspace_id":wid,"format":"json"})
    assert raw.status_code==200 and raw.json()["manifest"]["posts"][0]["caption"]==caption
    assert all(post["state"]=="approved" and post["visual"]["sha256"] for post in raw.json()["manifest"]["posts"])
    csv_response=authenticated_client.get(path,params={"workspace_id":wid,"format":"csv"})
    rows=list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert len(rows)==7 and rows[0]["caption"].startswith("'=") and rows[0]["visual_file"].startswith("creatives/")
    excel=authenticated_client.get(path,params={"workspace_id":wid,"format":"xlsx"})
    book=load_workbook(io.BytesIO(excel.content));assert book.active.max_row==8
    assert book.active["D2"].data_type=="s" and book.active["D2"].value==caption
    packed=authenticated_client.get(path,params={"workspace_id":wid,"format":"zip"})
    with zipfile.ZipFile(io.BytesIO(packed.content)) as archive:
        assert {"campaign.json","posts.csv","posts.xlsx"}<=set(archive.namelist())
        assert len([name for name in archive.namelist() if name.startswith("creatives/")])==7
    assert authenticated_client.get(path,params={"workspace_id":"foreign"}).status_code==404
    board=authenticated_client.get(base+"/board",params={"workspace_id":wid}).json()
    review=next(row for row in board if row["target_kind"]=="campaign_item" and row["target_id"]==iid)
    assert authenticated_client.post(f"/api/v1/approvals/{review['id']}/decide",params={"workspace_id":wid},json={"decision":"reopen"}).status_code==200
    assert not authenticated_client.get(base+"/finalized",params={"workspace_id":wid}).json()["current"]
    assert authenticated_client.get(path,params={"workspace_id":wid}).status_code==409


def test_incomplete_final_review_and_running_batch_do_not_execute(authenticated_client,monkeypatch):
    wid,plan=campaign(authenticated_client);base=f"/api/v1/campaigns/{plan['id']}"
    assert authenticated_client.post(base+"/finalize",params={"workspace_id":wid},json={"request_id":str(uuid4()),"reviewed":True}).status_code==409
    first=authenticated_client.post(base+"/visual-batches",params={"workspace_id":wid},json={"request_id":str(uuid4()),"generation_mode":"sample"}).json()
    with SessionLocal() as db:db.get(WorkflowRun,first["id"]).status="running";db.commit()
    monkeypatch.setattr(creatives.MetaClient,"image",lambda *args:pytest.fail("No replay of running batch"))
    assert authenticated_client.post(base+f"/visual-batches/{first['id']}/next",params={"workspace_id":wid}).status_code==409


def test_live_bulk_uses_one_image_per_step_with_shared_run_budget(authenticated_client,monkeypatch):
    wid,plan=campaign(authenticated_client);enable(monkeypatch)
    monkeypatch.setenv("META_RUN_BUDGET_CENTS","5");monkeypatch.setenv("META_WORKSPACE_BUDGET_CENTS","50")
    calls=[]
    def handle(request):
        assert request.url.path.endswith("/images/generations");calls.append(request)
        return httpx.Response(200,json={"created":1,"data":[{"b64_json":base64.b64encode(png()).decode()}]})
    monkeypatch.setattr(meta_client,"OpenAI",lambda **kwargs:SDK(**kwargs,http_client=httpx.Client(transport=httpx.MockTransport(handle))))
    base=f"/api/v1/campaigns/{plan['id']}"
    first=authenticated_client.post(base+"/visual-batches",params={"workspace_id":wid},json={"request_id":str(uuid4()),"generation_mode":"live","select_drafts":True}).json()
    path=base+f"/visual-batches/{first['id']}/next"
    assert authenticated_client.post(path,params={"workspace_id":wid}).json()["status"]=="pending"
    second=authenticated_client.post(path,params={"workspace_id":wid}).json()
    assert second["status"]=="failed" and len(second["completed"])==1 and len(calls)==1
    assert authenticated_client.post(path,params={"workspace_id":wid}).json()["status"]=="failed" and len(calls)==1
    assert len(authenticated_client.get(base+"/creative-status",params={"workspace_id":wid}).json())==1


def test_editable_image_prompt_uses_real_client_field_and_idempotent_paid_request(authenticated_client,monkeypatch):
    wid,plan=campaign(authenticated_client);enable(monkeypatch);calls=[]
    def handle(request):
        calls.append(json.loads(request.content));return httpx.Response(200,json={"created":1,"data":[{"b64_json":base64.b64encode(png()).decode()}]})
    monkeypatch.setattr(meta_client,"OpenAI",lambda **kwargs:SDK(**kwargs,http_client=httpx.Client(transport=httpx.MockTransport(handle))))
    route=f"/api/v1/campaigns/{plan['id']}/items/{plan['items'][0]['id']}/creatives"
    prompt="Use the owner-selected image direction. "+"wood "*170+"Preserve this final direction."
    body={"request_id":str(uuid4()),"variants":1,"evaluate":False,"prompt_override":prompt,"generation_mode":"live"}
    first=authenticated_client.post(route,params={"workspace_id":wid},json=body)
    assert first.status_code==201 and first.json()["created_count"]==1,first.text
    assert prompt in calls[0]["prompt"]
    assert authenticated_client.post(route,params={"workspace_id":wid},json=body).json()["created_count"]==1 and len(calls)==1
    assert authenticated_client.post(route,params={"workspace_id":wid},json={**body,"prompt_override":"A changed image direction for this owner."}).status_code==409


def test_optional_vision_failure_preserves_paid_image_and_does_not_replay(authenticated_client,monkeypatch):
    wid,plan=campaign(authenticated_client);enable(monkeypatch);calls=[]
    def handle(request):
        calls.append(request.url.path)
        if request.url.path.endswith("/images/generations"):
            return httpx.Response(200,json={"created":1,"data":[{"b64_json":base64.b64encode(png()).decode()}]})
        return httpx.Response(500,json={"error":{"message":"private provider error"}})
    monkeypatch.setattr(meta_client,"OpenAI",lambda **kwargs:SDK(**kwargs,http_client=httpx.Client(transport=httpx.MockTransport(handle))))
    route=f"/api/v1/campaigns/{plan['id']}/items/{plan['items'][0]['id']}/creatives"
    body={"request_id":str(uuid4()),"variants":1,"evaluate":True,"generation_mode":"live"}
    response=authenticated_client.post(route,params={"workspace_id":wid},json=body)
    assert response.status_code==201 and response.json()["created_count"]==1 and "Image saved" in response.json()["warning"]
    visual=response.json()["variants"][0]
    raw=authenticated_client.get(route+f"/{visual['id']}/image",params={"workspace_id":wid})
    assert raw.status_code==200 and raw.content==png()
    assert authenticated_client.post(route,params={"workspace_id":wid},json=body).json()["created_count"]==1 and len(calls)==2
